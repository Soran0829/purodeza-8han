"""Export the incident ledger to an Excel workbook for browsing."""

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = ["id", "date", "tag", "category", "title", "location", "reportedLocation",
          "locationAccuracy", "lat", "lng", "description", "body", "sourceUrl", "lastChecked"]
WIDTHS = {"id": 22, "date": 26, "tag": 20, "category": 18, "title": 34, "location": 30,
          "reportedLocation": 22, "locationAccuracy": 16, "lat": 12, "lng": 12,
          "description": 34, "body": 60, "sourceUrl": 52, "lastChecked": 22}
NUMERIC_FIELDS = {"lat", "lng"}


def load_rows(path):
    if path.suffix.lower() == ".json":
        return json.loads(path.read_text(encoding="utf-8"))
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def write_sheet(rows, path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "incidents"
    sheet.append(FIELDS)
    header_fill = PatternFill("solid", fgColor="101A33")
    for column, field in enumerate(FIELDS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = WIDTHS.get(field, 16)

    for row in rows:
        values = []
        for field in FIELDS:
            value = row.get(field, "")
            if field in NUMERIC_FIELDS:
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    value = ""
            values.append(value)
        sheet.append(values)
        url_cell = sheet.cell(row=sheet.max_row, column=FIELDS.index("sourceUrl") + 1)
        if url_cell.value:
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column_letter in {"K", "L"}:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif cell.column_letter == "F":
                cell.alignment = Alignment(vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description="不審者情報の台帳をExcelファイルに変換")
    parser.add_argument("--source", type=Path, default=Path("data/incidents.json"),
                        help="入力ファイル（.json または .csv、既定は data/incidents.json）")
    parser.add_argument("--xlsx", type=Path, default=Path("data/patnet_aichi_archive.xlsx"))
    args = parser.parse_args()
    if not args.source.exists():
        raise SystemExit(f"入力ファイルが見つかりません: {args.source}")
    rows = load_rows(args.source)
    write_sheet(rows, args.xlsx)
    print(f"Excel出力: {args.xlsx} / {len(rows)}件")


if __name__ == "__main__":
    main()
